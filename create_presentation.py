from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn


PRIMARY_BLUE = RGBColor(20, 90, 160)
ACCENT_BLUE = RGBColor(35, 140, 200)
TEXT_DARK = RGBColor(34, 34, 34)
TEXT_GRAY = RGBColor(102, 102, 102)


def set_background(prs, slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 248, 252)

    top_bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0),
        Inches(0),
        prs.slide_width,
        Inches(0.5),
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = PRIMARY_BLUE
    top_bar.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        prs.slide_width - Inches(3),
        Inches(0.15),
        Inches(3),
        Inches(0.2),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_BLUE
    accent.line.fill.background()


def apply_transition(slide, effect="fade"):
    sld = slide._element
    existing = sld.find(qn("p:transition"))
    if existing is not None:
        sld.remove(existing)

    transition = OxmlElement("p:transition")
    transition.set("spd", "fast")
    transition.set("type", effect)
    sld.append(transition)


def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    set_background(prs, slide)

    title_box = slide.shapes.title
    title_box.text = title
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.font.size = Pt(40)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = RGBColor(255, 255, 255)

    subtitle_box = slide.placeholders[1]
    subtitle_box.text = subtitle
    subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
    subtitle_paragraph.font.size = Pt(20)
    subtitle_paragraph.font.color.rgb = RGBColor(230, 230, 230)

    apply_transition(slide, effect="cut")


def add_bullet_slide(prs, title, bullets):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    set_background(prs, slide)

    title_box = slide.shapes.title
    title_box.text = title
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.font.size = Pt(30)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PRIMARY_BLUE

    content = slide.shapes.placeholders[1]
    text_frame = content.text_frame
    text_frame.clear()
    text_frame.word_wrap = True

    for idx, bullet in enumerate(bullets):
        p = text_frame.add_paragraph() if idx > 0 else text_frame.paragraphs[0]
        p.text = bullet
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT_DARK
        p.level = 0

    apply_transition(slide, effect="fade")


def add_table_slide(prs, title, table_data):
    slide_layout = prs.slide_layouts[5]  # Title only
    slide = prs.slides.add_slide(slide_layout)
    set_background(prs, slide)

    title_box = slide.shapes.title
    title_box.text = title
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.font.size = Pt(30)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PRIMARY_BLUE

    rows = len(table_data)
    cols = len(table_data[0])

    table = slide.shapes.add_table(rows, cols, Inches(0.7), Inches(1.7), Inches(11.3), Inches(4.8)).table

    for c in range(cols):
        column = table.columns[c]
        column.width = Inches(2.8 if c == 0 else 2.5)

    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = table_data[r][c]
            paragraphs = cell.text_frame.paragraphs
            for paragraph in paragraphs:
                paragraph.font.size = Pt(18 if r == 0 else 16)
                paragraph.font.bold = r == 0
                paragraph.font.color.rgb = PRIMARY_BLUE if r == 0 else TEXT_DARK
            cell.vertical_anchor = 1

    for cell in table.rows[0].cells:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(230, 240, 250)

    apply_transition(slide, effect="push")


def add_summary_slide(prs, title, highlight, bullets):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    set_background(prs, slide)

    title_box = slide.shapes.title
    title_box.text = title
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.font.size = Pt(30)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PRIMARY_BLUE

    body_shape = slide.shapes.placeholders[1]
    text_frame = body_shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True

    highlight_paragraph = text_frame.add_paragraph()
    highlight_paragraph.text = highlight
    highlight_paragraph.font.size = Pt(22)
    highlight_paragraph.font.bold = True
    highlight_paragraph.font.color.rgb = PRIMARY_BLUE

    for bullet in bullets:
        p = text_frame.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_GRAY
        p.level = 1

    apply_transition(slide, effect="wipe")


def build_presentation():
    prs = Presentation()

    add_title_slide(
        prs,
        "2024年度医疗项目经理述职报告",
        "项目经理：张三 | 日期：2025年1月",
    )

    add_bullet_slide(
        prs,
        "年度工作概览",
        [
            "负责智慧病区升级、远程会诊平台、移动护理等 5 个重点项目",
            "覆盖 12 家三甲医院与区域医疗中心，项目合同总额约 4.8 亿元",
            "实现项目集群 PMO 管理，建立统一的进度、成本与质量监控机制",
        ],
    )

    add_bullet_slide(
        prs,
        "年度核心成果",
        [
            "智慧病区项目全量上线，床旁移动护理工作效率提升 32%",
            "远程会诊平台完成 1860 场跨区域病例讨论，专家响应时间缩短 40%",
            "区域质控数据共享率提升至 95%，辅助决策模型上线 3 个临床科室",
        ],
    )

    add_table_slide(
        prs,
        "关键指标达成情况",
        [
            ["指标", "年度目标", "实际完成", "完成率"],
            ["项目按期交付率", "≥ 90%", "94.6%", "105%"],
            ["客户满意度", ">= 4.5/5", "4.7/5", "104%"],
            ["预算控制偏差", "±5%", "+2.8%", "合规"],
            ["重大风险闭环率", "100%", "100%", "达成"],
        ],
    )

    add_bullet_slide(
        prs,
        "项目管理亮点",
        [
            "建立周度里程碑评审机制，需求变更平均响应时间缩短至 1.5 天",
            "应用 Earned Value 管理模型，实现跨项目进度可视化",
            "推行数字化验收流程，验收效率提升 25%",
        ],
    )

    add_bullet_slide(
        prs,
        "团队与资源协同",
        [
            "组建 18 人跨职能项目团队，覆盖临床、信息、运维三大模块",
            "开展 12 场内部培训与案例复盘，培养 4 名预备项目经理",
            "签约 6 家战略供应商，建立绩效共担激励机制",
        ],
    )

    add_bullet_slide(
        prs,
        "风险控制与合规管理",
        [
            "风险热力图动态更新，识别 27 项较大风险并完成闭环处理",
            "推进网络安全等保三级测评，通过率 100%",
            "联合法务与合规团队完成 8 次合同审查与数据合规检查",
        ],
    )

    add_bullet_slide(
        prs,
        "客户与多方协同",
        [
            "建立 C-level 共创委员会，每季度推进战略需求对齐",
            "引入临床科室联络官制度，满意度调查覆盖率 100%",
            "与政府卫健委、医保局保持季度例会，保障政策对接",
        ],
    )

    add_bullet_slide(
        prs,
        "问题与改进方向",
        [
            "部分项目基础数据治理仍不完善，计划引入主数据平台建设",
            "跨院区部署资源紧张，拟建设统一的云资源池与自动化运维平台",
            "临床用户培训需要提前规划，拟增设面向护士的混合式培训体系",
        ],
    )

    add_summary_slide(
        prs,
        "2025 年工作规划",
        "围绕“质量、创新、协同”三条主线推进项目群升级",
        [
            "质量：建立端到端交付标准，推进全流程数据驱动管理",
            "创新：落地 AI 辅助诊疗试点项目，实现 3 个可复制样板",
            "协同：打造区域项目指挥中心，实现跨院区资源共享",
        ],
    )

    add_bullet_slide(
        prs,
        "感谢与展望",
        [
            "感谢领导与同事的支持，团队将持续聚焦临床价值创造",
            "期待 2025 年与更多合作伙伴共建智慧医疗标杆",
        ],
    )

    return prs


def main():
    presentation = build_presentation()
    output_path = "医疗项目经理述职报告.pptx"
    presentation.save(output_path)
    print(f"已生成 {output_path}")


if __name__ == "__main__":
    main()
